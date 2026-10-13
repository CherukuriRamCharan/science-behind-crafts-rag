
import json
import hashlib
from pathlib import Path

import pandas as pd
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data" / "raw"
OUTPUT = ROOT / "data" / "extracted" / "documents.jsonl"

SUPPORTED = {".pdf", ".docx", ".xlsx", ".xlsm"}


def sha256_file(path):
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def extract_pdf(path):
    reader = PdfReader(str(path))
    parts = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(f"[Page {page_number}]\n{text}")

    return "\n\n".join(parts), len(reader.pages)


def extract_docx(path):
    document = Document(str(path))
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        parts.append(f"[Table {table_index}]")
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(values))

    return "\n\n".join(parts), None


def extract_excel(path):
    parts = []
    workbook = pd.ExcelFile(path)

    for sheet_name in workbook.sheet_names:
        frame = pd.read_excel(
            workbook, sheet_name=sheet_name, dtype=object
        )
        parts.append(f"[Sheet: {sheet_name}]")

        if frame.empty:
            parts.append("[Empty sheet]")
            continue

        frame = frame.fillna("")
        columns = [str(column) for column in frame.columns]
        parts.append("Columns: " + " | ".join(columns))

        for row_number, row in enumerate(
            frame.itertuples(index=False, name=None), start=1
        ):
            values = [
                f"{column}: {value}"
                for column, value in zip(columns, row)
                if str(value).strip()
            ]
            if values:
                parts.append(
                    f"[Row {row_number}] " + " | ".join(values)
                )

    return "\n\n".join(parts), None


def extract_file(path):
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return extract_excel(path)

    raise ValueError(f"Unsupported file type: {suffix}")


def main():
    if not RAW.exists():
        raise FileNotFoundError(f"Raw folder does not exist: {RAW}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    errors = []

    with OUTPUT.open("w", encoding="utf-8") as out:
        for path in sorted(RAW.rglob("*")):
            if not path.is_file() or path.suffix.lower() not in SUPPORTED:
                continue

            try:
                text, page_count = extract_file(path)

                record = {
                    "document_key": sha256_file(path),
                    "file_name": path.name,
                    "relative_path": str(path.relative_to(ROOT)),
                    "file_type": path.suffix.lower(),
                    "page_count": page_count,
                    "text": text,
                    "extraction_status": (
                        "ok" if text.strip() else "empty"
                    ),
                }

                out.write(
                    json.dumps(record, ensure_ascii=False) + "\n"
                )
                count += 1

                print(
                    f"[{record['extraction_status']}] {path.name}"
                )

            except Exception as exc:
                errors.append({
                    "path": str(path),
                    "error": str(exc),
                })
                print(f"[ERROR] {path}: {exc}")

    error_path = OUTPUT.parent / "extraction_errors.json"
    error_path.write_text(
        json.dumps(errors, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(f"\nDocuments extracted: {count}")
    print(f"Errors: {len(errors)}")
    print(f"Output: {OUTPUT}")
    print(f"Error log: {error_path}")


if __name__ == "__main__":
    main()
